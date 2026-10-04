"""Validate actual exported workbook styles, score rules and editable formula totals."""
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

path='/tmp/vl2-records.xlsx'
with ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith(('.xml','.rels')): ET.fromstring(z.read(name))
    assert any(n.startswith('xl/media/') for n in z.namelist())
w=load_workbook(path)
assert len(w.sheetnames)==6
answers=w['學生探究答案']; scores=w['教師評分']; data=w['量度計算與圖點']
assert answers['B2'].value=='陳小明'
assert len(w['裝置設計圖']._images)==1
assert answers['M2'].font.color.rgb=='FF00834A'  # correct independent variable
assert answers['S2'].font.color.rgb=='FFC03030'  # deliberately incorrect assumptions
assert answers['H2'].font.color.rgb=='FF173E34'  # open observation remains unmarked
assert len({answers[c+'2'].fill.fgColor.rgb for c in ['H','M','S','X','Z','AG']})==6
assert any(cell.font.color.rgb=='FFC03030' for row in data.iter_rows(min_row=2) for cell in row)
assert len(scores.data_validations.dataValidation)==9
assert len(answers.conditional_formatting)==8
assert len(scores.conditional_formatting)==27
assert w.calculation.fullCalcOnLoad and w.calculation.forceFullCalc
headers={cell.value:cell.column for cell in scores[1]}
manual=[scores[v.sqref.__str__().split(':')[0]] for v in scores.data_validations.dataValidation]
assert all(c.value=='' or c.value is None for c in manual)
assert len(w['評分準則']['A'])>=19

# Evaluate the small subset of Excel functions used by this workbook, using real
# exported formulas and referenced cells; distinguish blanks from explicit zero.
def number(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def ev(cell,seen=None):
    seen=set() if seen is None else seen
    assert cell not in seen, 'circular formula'
    value=scores[cell].value
    if not isinstance(value,str) or not value.startswith('='): return '' if value is None else value
    seen=seen|{cell}
    # Replace references only outside quoted strings.
    parts=value[1:].split('"')
    expression=''.join(('"'+part+'"') if i%2 else re.sub(r'\b[A-Z]+[0-9]+\b',lambda m:repr(ev(m[0],seen)),part) for i,part in enumerate(parts))
    expression=re.sub(r'(?<![<>=!])=(?!=)','==',expression)
    funcs={'IF':lambda cond,a,b:a if cond else b,'AND':lambda *xs:all(xs),'COUNT':lambda *xs:sum(number(x) for x in xs),'SUM':lambda *xs:sum(x for x in xs if number(x)),'ROUND':lambda x,n:round(x,n)}
    return eval(expression,{'__builtins__':{}},funcs)
assert ev('AE2')=='待評／未完成'  # overall score
assert ev('Y2')=='待評'           # SPS subtotal
assert ev('AD2')=='待評'          # knowledge subtotal
assert ev('G2')=='待評'           # observation requires teacher score
assert ev('K2')==4               # correct variable choices
assert ev('M2')==0               # incorrect assumption set
assert ev('S2')==1.5             # D calculation intentionally wrong
assert ev('T2')==2               # correct conclusions
assert ev('V2')==2               # D point matches its own wrong calculation
for validation in scores.data_validations.dataValidation:
    cell=str(validation.sqref).split(':')[0]; scores[cell]=int(validation.formula2)
assert ev('AD2')==8
sps=sum(ev(c+'2') for c in ['G','K','O','R','U','X'])
assert ev('Y2')==round(sps,2)
assert ev('AE2')==round(sps+8,2)
assert ev('AF2')=='評分完成'
scores['Z2']=0;assert ev('AD2')==6
scores['Z2']=None;assert ev('AD2')=='待評';assert ev('AE2')=='待評／未完成'
scores['Z2']=2;scores['D2']='待提交反思';assert ev('AE2')=='待評／未完成'
# Actual exported conditional rules must follow the manual score, preserving fill.
def colour_matches(target):
    rules=next(rules for cf,rules in answers.conditional_formatting._cf_rules.items() if str(cf.sqref)==target)
    matches=[]
    for rule in rules:
        formula=rule.formula[0]
        refs=re.findall(r"!([A-Z]+[0-9]+)",formula)
        if any(not number(ev(ref)) for ref in refs): continue
        formula=re.sub(r'INDIRECT\("\'教師評分\'!([A-Z]+[0-9]+)"\)',lambda m:repr(ev(m[1])),formula)
        formula=re.sub(r'(?<![<>=!])=(?!=)','==',formula)
        if eval(formula,{'__builtins__':{}},{'AND':lambda *xs:all(xs),'ISNUMBER':number}): matches.append(rule.dxfId)
    return matches
scores['E2']=None;assert colour_matches('H2')==[]
scores['E2']=2;assert colour_matches('H2')==[0]  # green full credit
scores['E2']=0;assert colour_matches('H2')==[1]  # red zero credit
scores['E2']=1;assert colour_matches('H2')==[2]  # amber partial credit
scores['Z2']=None;assert colour_matches('AG2')==[]
scores['Z2']=2;assert colour_matches('AG2')==[0]
for cell in ['Z2','AA2','AB2','AC2']: scores[cell]=0
assert colour_matches('AG2')==[1]
print('PASS: XLSX colours, category fills, nine validated manual fields, formulas, pending vs zero, six SPS totals, reflection knowledge score, overall score, conditional colours, images.')
